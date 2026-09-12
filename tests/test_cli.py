"""End to end: the command line runs every stage on a synthetic video with the stand-in model."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from specto.cli import main
from specto.model import Analysis, Recording

VTT = """WEBVTT

00:00:00.000 --> 00:00:02.500
<v Sam>This is the customer search screen, we type the postcode here.

00:00:03.200 --> 00:00:05.800
<v Sam>Then we open the customer details and check the account number.

00:00:06.100 --> 00:00:08.700
<v Sam>Save sends it to the approval queue, only a team lead can approve.

00:00:09.300 --> 00:00:11.500
<v Sam>And that is it, the customer is done.
"""


def test_cli_run_with_fake_model(synthetic_video: Path, tmp_path: Path) -> None:
    vtt = tmp_path / "walkthrough.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    out = tmp_path / "out"

    rc = main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(out), "--fake"])
    assert rc == 0

    for name in ("recording.json", "analysis.json", "analysis.xlsx", "report.md"):
        assert (out / name).exists(), name
    recording = Recording.model_validate_json((out / "recording.json").read_text())
    analysis = Analysis.model_validate_json((out / "analysis.json").read_text())
    assert recording.transcript_source == "file"
    assert len(recording.segments) == 4
    assert analysis.requirements and analysis.questions
    for kf in recording.keyframes:
        assert (out / kf.path).exists()

    wb = openpyxl.load_workbook(out / "analysis.xlsx")
    assert wb.sheetnames[0] == "Summary"
    assert "SME Questions" in wb.sheetnames

    # second run reuses the saved stages and still exports
    rc = main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(out), "--fake"])
    assert rc == 0

    # export alone rebuilds the workbook from the saved analysis
    (out / "analysis.xlsx").unlink()
    assert main(["export", str(out)]) == 0
    assert (out / "analysis.xlsx").exists()


def test_cli_refuses_without_key(synthetic_video: Path, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    vtt = tmp_path / "w.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    rc = main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(tmp_path / "o")])
    assert rc == 2


def test_cli_estimate_stops_before_the_model(synthetic_video: Path, tmp_path: Path, capsys) -> None:
    vtt = tmp_path / "w.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    out = tmp_path / "o"
    assert main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(out), "--estimate"]) == 0
    text = capsys.readouterr().out
    assert "Estimated cost on claude-opus-5" in text
    assert "claude-haiku-4-5" in text
    assert not (out / "analysis.json").exists()


def test_cli_doctor_runs(capsys) -> None:
    rc = main(["doctor"])
    assert rc in (0, 1)
    assert "ffmpeg" in capsys.readouterr().out


def test_cli_live_replay_with_fake_model(tmp_path: Path) -> None:
    from PIL import Image, ImageDraw

    shots = tmp_path / "shots"
    shots.mkdir()
    for t, colour, label in ((0, (40, 70, 140), "Search"), (12, (245, 245, 245), "Details"), (30, (30, 120, 60), "Queue")):
        img = Image.new("RGB", (640, 360), colour)
        ImageDraw.Draw(img).text((40, 40), label, fill=(0, 0, 0) if sum(colour) > 380 else (255, 255, 255))
        img.save(shots / f"shot_{t}.png")
    vtt = tmp_path / "w.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    out = tmp_path / "call"
    rc = main(["live", "--out", str(out), "--replay", str(shots), "--transcript", str(vtt), "--every", "20", "--fake"])
    assert rc == 0
    assert (out / "live_questions.md").exists()
    assert (out / "analysis.xlsx").exists()


def test_cli_merge_two_runs(synthetic_video: Path, tmp_path: Path) -> None:
    vtt = tmp_path / "w.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    a, b = tmp_path / "a", tmp_path / "b"
    assert main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(a), "--fake"]) == 0
    assert main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(b), "--fake"]) == 0
    merged = tmp_path / "all"
    assert main(["merge", str(a), str(b), "--out", str(merged)]) == 0
    rec = Recording.model_validate_json((merged / "recording.json").read_text())
    ana = Analysis.model_validate_json((merged / "analysis.json").read_text())
    assert len(rec.keyframes) == 2 * len(Recording.model_validate_json((a / "recording.json").read_text()).keyframes)
    assert "merged" in ana.title
    assert (merged / "analysis.xlsx").exists() and (merged / "report.html").exists()


def test_cli_answers_then_resolve_with_fake_model(synthetic_video: Path, tmp_path: Path) -> None:
    import openpyxl

    vtt = tmp_path / "w.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    out = tmp_path / "o"
    assert main(["run", str(synthetic_video), "--transcript", str(vtt), "--out", str(out), "--fake"]) == 0
    before = Analysis.model_validate_json((out / "analysis.json").read_text())
    first_q = before.questions[0].id

    wb = openpyxl.load_workbook(out / "analysis.xlsx")
    ws = wb["SME Questions"]
    headers = [c.value for c in ws[1]]
    id_col, ans_col = headers.index("Id") + 1, headers.index("Answer") + 1
    for row in ws.iter_rows(min_row=2):
        if row[id_col - 1].value == first_q:
            row[ans_col - 1].value = "Only team leads, and it must be logged."
    wb.save(out / "analysis.xlsx")

    assert main(["answers", str(out)]) == 0
    after = Analysis.model_validate_json((out / "analysis.json").read_text())
    answered = next(q for q in after.questions if q.id == first_q)
    assert answered.status == "answered" and "team leads" in (answered.answer or "")
    html = (out / "report.html").read_text()
    assert "Only team leads, and it must be logged." in html

    assert main(["resolve", str(out), "--fake"]) == 0
    resolved = Analysis.model_validate_json((out / "analysis.json").read_text())
    assert len(resolved.requirements) > len(before.requirements)
