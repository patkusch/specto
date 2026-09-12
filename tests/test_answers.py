"""Reading the analyst's answers back out of the workbook. No API, no network."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from specto.answers import ImportResult, apply_answers, import_answers, normalise_status, read_answers
from specto.export import export_xlsx
from specto.model import Analysis, Recording

FIXTURES = Path(__file__).parent / "fixtures"

Q002_ANSWER = "No. The email is optional. The record saves without one."
Q003_ANSWER = "Passport or driving licence as PDF or JPEG, up to 10 MB."


@pytest.fixture
def analysis() -> Analysis:
    return Analysis.model_validate(json.loads((FIXTURES / "sample_analysis.json").read_text()))


@pytest.fixture
def recording() -> Recording:
    return Recording.model_validate(json.loads((FIXTURES / "sample_recording.json").read_text()))


def columns(ws) -> dict[str, int]:
    """Header text -> 1-based column number, so tests never assume positions."""
    return {str(c.value): c.column for c in ws[1] if c.value}


def rows_by_id(ws) -> dict[str, int]:
    col = columns(ws)["Id"]
    return {ws.cell(row=r, column=col).value: r for r in range(2, ws.max_row + 1)}


def type_answers(xlsx: Path, entries: dict[str, tuple[str | None, str | None]]) -> None:
    """Type (answer, status) into the SME Questions sheet the way an analyst would."""
    wb = load_workbook(xlsx)
    ws = wb["SME Questions"]
    col = columns(ws)
    rows = rows_by_id(ws)
    for question_id, (answer, status) in entries.items():
        if question_id not in rows:
            ws.append([question_id])
            rows[question_id] = ws.max_row
        r = rows[question_id]
        if answer is not None:
            ws.cell(row=r, column=col["Answer"]).value = answer
        if status is not None:
            ws.cell(row=r, column=col["Status"]).value = status
    wb.save(xlsx)


@pytest.fixture
def answered_workbook(analysis, recording, tmp_path) -> Path:
    """The exported workbook with two answers typed and one question dropped."""
    xlsx = export_xlsx(analysis, recording, tmp_path)
    type_answers(xlsx, {
        "Q002": (Q002_ANSWER, None),
        "Q003": (Q003_ANSWER, "Done"),
        "Q004": (None, "n/a"),
    })
    return xlsx


def test_read_answers_returns_only_touched_rows(answered_workbook):
    assert read_answers(answered_workbook) == {
        "Q002": (Q002_ANSWER, "answered"),
        "Q003": (Q003_ANSWER, "answered"),
        "Q004": (None, "not needed"),
    }


def test_apply_answers_sets_answer_and_status(analysis, answered_workbook):
    analysis, result = apply_answers(analysis, read_answers(answered_workbook))

    assert result == ImportResult(changed=3, unknown_ids=[])
    by_id = {q.id: q for q in analysis.questions}
    assert by_id["Q001"].status == "open" and by_id["Q001"].answer is None
    assert by_id["Q002"].status == "answered" and by_id["Q002"].answer == Q002_ANSWER
    assert by_id["Q003"].status == "answered" and by_id["Q003"].answer == Q003_ANSWER
    assert by_id["Q004"].status == "not needed" and by_id["Q004"].answer is None

    # Applying the same sheet again changes nothing.
    _, again = apply_answers(analysis, read_answers(answered_workbook))
    assert again.changed == 0


def test_import_answers_writes_analysis_json_and_nothing_else(analysis, answered_workbook, tmp_path):
    (tmp_path / "analysis.json").write_text(analysis.model_dump_json(indent=2))
    before = {p.name: p.stat().st_mtime_ns for p in tmp_path.iterdir()}

    result = import_answers(tmp_path)

    assert result.changed == 3
    assert result.unknown_ids == []
    reloaded = Analysis.model_validate_json((tmp_path / "analysis.json").read_text())
    by_id = {q.id: q for q in reloaded.questions}
    assert by_id["Q002"].answer == Q002_ANSWER
    assert by_id["Q003"].status == "answered"
    assert by_id["Q004"].status == "not needed"
    # The workbook was read, not rewritten.
    assert (tmp_path / "analysis.xlsx").stat().st_mtime_ns == before["analysis.xlsx"]


def test_import_answers_from_a_copy_elsewhere(analysis, answered_workbook, tmp_path):
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    (out_dir / "analysis.json").write_text(analysis.model_dump_json(indent=2))
    copy = tmp_path / "analysis (Dave's notes).xlsx"
    copy.write_bytes(answered_workbook.read_bytes())

    result = import_answers(out_dir, xlsx_path=copy)

    assert result.changed == 3
    reloaded = Analysis.model_validate_json((out_dir / "analysis.json").read_text())
    assert {q.id: q.status for q in reloaded.questions} == {
        "Q001": "open", "Q002": "answered", "Q003": "answered", "Q004": "not needed",
    }


def test_columns_found_by_header_after_a_column_is_inserted(answered_workbook):
    wb = load_workbook(answered_workbook)
    ws = wb["SME Questions"]
    ws.insert_cols(1)
    ws.cell(row=1, column=1).value = "Notes to self"
    ws.cell(row=2, column=1).value = "check with Dave"
    wb.save(answered_workbook)

    assert read_answers(answered_workbook) == {
        "Q002": (Q002_ANSWER, "answered"),
        "Q003": (Q003_ANSWER, "answered"),
        "Q004": (None, "not needed"),
    }


def test_unknown_id_is_reported_not_applied(analysis, answered_workbook):
    type_answers(answered_workbook, {"Q999": ("A question that is not in the analysis.", None)})

    answers = read_answers(answered_workbook)
    assert answers["Q999"] == ("A question that is not in the analysis.", "answered")

    analysis, result = apply_answers(analysis, answers)
    assert result.changed == 3
    assert result.unknown_ids == ["Q999"]
    assert [q.id for q in analysis.questions] == ["Q001", "Q002", "Q003", "Q004"]


def test_answer_with_open_status_keeps_the_question_open(analysis, recording, tmp_path):
    xlsx = export_xlsx(analysis, recording, tmp_path)
    type_answers(xlsx, {"Q001": ("Dave thinks it warns. To confirm.", "open")})

    assert read_answers(xlsx) == {"Q001": ("Dave thinks it warns. To confirm.", None)}
    analysis, result = apply_answers(analysis, read_answers(xlsx))
    assert result.changed == 1
    assert analysis.questions[0].status == "open"
    assert analysis.questions[0].answer == "Dave thinks it warns. To confirm."


def test_missing_sheet_or_header_is_a_clear_error(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Other"
    wb.save(tmp_path / "no_sheet.xlsx")
    with pytest.raises(ValueError, match="SME Questions"):
        read_answers(tmp_path / "no_sheet.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "SME Questions"
    ws.append(["Id", "Question", "Answer"])
    wb.save(tmp_path / "no_status.xlsx")
    with pytest.raises(ValueError, match="Status"):
        read_answers(tmp_path / "no_status.xlsx")


@pytest.mark.parametrize("typed, expected", [
    ("answered", "answered"),
    ("Answered", "answered"),
    ("  DONE ", "answered"),
    ("yes", "answered"),
    ("closed", "answered"),
    ("Resolved", "answered"),
    ("not needed", "not needed"),
    ("Not  Needed", "not needed"),
    ("n/a", "not needed"),
    ("NA", "not needed"),
    ("drop", "not needed"),
    ("no longer needed", "not needed"),
    ("open", None),
    ("", None),
    ("   ", None),
    (None, None),
    ("waiting on Dave", None),
])
def test_status_normalisation(typed, expected):
    assert normalise_status(typed) == expected
