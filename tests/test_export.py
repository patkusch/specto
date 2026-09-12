"""Stage 3 tests: the workbook and the report come out with the right shape and links."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from specto.export import export_all, export_markdown, export_xlsx
from specto.model import Analysis, Recording
from specto.timefmt import mmss

FIXTURES = Path(__file__).parent / "fixtures"

SHEET_ORDER = [
    "Summary", "Journey", "Screens", "Data Fields", "Actions",
    "Requirements", "Acceptance Criteria", "SME Questions", "Transcript",
]

HEADERS = {
    "Summary": ["Item", "Value"],
    "Journey": ["Order", "Screen", "What happens", "Who", "Time", "Frame"],
    "Screens": ["Id", "Name", "Purpose", "First seen", "Frame", "Frame indexes", "Fields", "Actions"],
    "Data Fields": ["Id", "Screen", "Label", "Type", "Required", "Example value", "Source", "Time", "Frame", "Notes"],
    "Actions": ["Id", "Screen", "Action", "Control", "Leads to", "Time", "Frame"],
    "Requirements": ["Id", "Statement", "Kind", "Priority", "Confidence", "Screen", "Rationale", "Source quote",
                     "Time", "Frame", "Acceptance criteria", "Writing check"],
    "Acceptance Criteria": ["Id", "Requirement id", "Requirement", "Given", "When", "Then", "Time", "Frame",
                            "Writing check"],
    "SME Questions": ["Id", "Question", "Why it matters", "Category", "Screen", "What was said", "Time", "Frame",
                      "Answer", "Status"],
    "Transcript": ["Time", "Speaker", "Text", "Frame"],
}


@pytest.fixture
def analysis() -> Analysis:
    return Analysis.model_validate(json.loads((FIXTURES / "sample_analysis.json").read_text()))


@pytest.fixture
def recording() -> Recording:
    return Recording.model_validate(json.loads((FIXTURES / "sample_recording.json").read_text()))


@pytest.fixture
def workbook(analysis, recording, tmp_path):
    path = export_xlsx(analysis, recording, tmp_path)
    assert path == tmp_path / "analysis.xlsx"
    assert path.exists()
    return load_workbook(path)


def header_row(ws) -> list:
    return [c.value for c in ws[1]]


def test_mmss():
    assert mmss(65) == "01:05"
    assert mmss(3725) == "1:02:05"
    assert mmss(0) == "00:00"
    assert mmss(59.9) == "00:59"


def test_sheet_names_in_order(workbook):
    assert workbook.sheetnames == SHEET_ORDER


def test_header_rows(workbook):
    for name, headers in HEADERS.items():
        assert header_row(workbook[name]) == headers, name


def test_row_counts_match_fixture(workbook, analysis, recording):
    expected = {
        "Journey": len(analysis.journey),
        "Screens": len(analysis.screens),
        "Data Fields": len(analysis.fields),
        "Actions": len(analysis.actions),
        "Requirements": len(analysis.requirements),
        "Acceptance Criteria": len(analysis.acceptance_criteria),
        "SME Questions": len(analysis.questions),
        "Transcript": len(recording.segments),
    }
    for name, count in expected.items():
        assert workbook[name].max_row - 1 == count, name


def test_summary_sheet_content(workbook, analysis, recording):
    ws = workbook["Summary"]
    items = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert items["Title"] == analysis.title
    assert items["Summary"] == analysis.summary
    assert items["Actors"] == "Onboarding clerk, Team manager"
    assert items["Recording"] == recording.source
    assert items["Duration"] == "05:00"
    assert items["Screens"] == len(analysis.screens)
    assert items["Data fields"] == len(analysis.fields)
    assert items["Requirements"] == len(analysis.requirements)
    assert items["Acceptance criteria"] == len(analysis.acceptance_criteria)
    assert items["Questions"] == len(analysis.questions)
    assert items["Model"] == analysis.usage.model
    assert items["Input tokens"] == analysis.usage.input_tokens
    assert "Generated" in items


def test_requirement_row_links_to_keyframe(workbook, analysis, recording):
    ws = workbook["Requirements"]
    headers = header_row(ws)
    frame_col = headers.index("Frame") + 1
    for row_number, req in enumerate(analysis.requirements, start=2):
        cell = ws.cell(row=row_number, column=frame_col)
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == recording.keyframes[req.keyframe_index].path
        assert cell.value == f"frame {req.keyframe_index} @ {mmss(req.timestamp)}"
        assert cell.font.underline == "single"
    first = ws.cell(row=2, column=frame_col)
    assert first.hyperlink.target == "frames/frame_0000.jpg"


def test_requirements_list_their_criteria(workbook):
    ws = workbook["Requirements"]
    col = header_row(ws).index("Acceptance criteria") + 1
    assert ws.cell(row=2, column=col).value == "AC001, AC002"
    assert ws.cell(row=4, column=col).value == "AC005"


def test_questions_sheet_has_answer_and_status_columns(workbook):
    ws = workbook["SME Questions"]
    headers = header_row(ws)
    assert headers[-2:] == ["Answer", "Status"]
    for row in ws.iter_rows(min_row=2, min_col=len(headers) - 1):
        assert all(c.value in (None, "") for c in row)


def test_freeze_panes_and_autofilter_on_every_sheet(workbook):
    for ws in workbook.worksheets:
        assert ws.freeze_panes == "A2", ws.title
        assert ws.auto_filter.ref, ws.title
        assert ws.auto_filter.ref.startswith("A1:"), ws.title


def test_header_style_and_column_widths(workbook):
    ws = workbook["Requirements"]
    for cell in ws[1]:
        assert cell.font.bold
        assert cell.fill.fill_type == "solid"
    widths = [ws.column_dimensions[c.column_letter].width for c in ws[1]]
    assert all(w is not None and 8 <= w <= 70 for w in widths)
    statement = ws.cell(row=2, column=2)
    assert statement.alignment.wrap_text
    assert statement.alignment.vertical == "top"


def test_transcript_rows_link_to_moment_frame(workbook, recording):
    ws = workbook["Transcript"]
    # First segment starts at 2s, inside moment 0 -> frame_0000; last starts at 262s, inside moment 5.
    assert ws.cell(row=2, column=4).hyperlink.target == "frames/frame_0000.jpg"
    assert ws.cell(row=ws.max_row, column=4).hyperlink.target == "frames/frame_0005.jpg"
    assert ws.cell(row=2, column=1).value == "00:02"
    assert ws.cell(row=2, column=2).value == "Expert"


def test_unknown_keyframe_written_without_link(analysis, recording, tmp_path):
    analysis.requirements[0].keyframe_index = 99
    path = export_xlsx(analysis, recording, tmp_path, filename="odd.xlsx")
    ws = load_workbook(path)["Requirements"]
    cell = ws.cell(row=2, column=header_row(ws).index("Frame") + 1)
    assert cell.hyperlink is None
    assert cell.value == "frame 99"


def test_markdown_report(analysis, recording, tmp_path):
    path = export_markdown(analysis, recording, tmp_path)
    assert path == tmp_path / "report.md"
    text = path.read_text()
    assert text.startswith(f"# {analysis.title}")
    for req in analysis.requirements:
        assert req.id in text
    for ac in analysis.acceptance_criteria:
        assert ac.id in text
    for q in analysis.questions:
        assert q.question in text
    for screen in analysis.screens:
        assert f"### {screen.id}: {screen.name}" in text
    assert "1. **Customer search**" in text
    assert "([frame 0 @ 00:15](frames/frame_0000.jpg))" in text
    assert "](frames/frame_0005.jpg)" in text


def test_export_all(analysis, recording, tmp_path):
    paths = export_all(analysis, recording, tmp_path)
    assert set(paths) == {"xlsx", "markdown", "html", "jira_csv", "ado_csv"}
    assert paths["xlsx"] == tmp_path / "analysis.xlsx"
    assert paths["markdown"] == tmp_path / "report.md"
    assert all(p.exists() for p in paths.values())


# --------------------------------------------------------------- writing check


def test_writing_check_column_on_requirements_and_criteria(workbook):
    for sheet in ("Requirements", "Acceptance Criteria"):
        assert header_row(workbook[sheet])[-1] == "Writing check", sheet


def test_dirty_requirement_row_carries_a_finding(workbook):
    ws = workbook["Requirements"]
    col = header_row(ws).index("Writing check") + 1
    by_id = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=col).value
             for r in range(2, ws.max_row + 1)}
    # R002 joins two rules with "or"; R003 is a well-formed user story.
    assert by_id["R002"].startswith('warn: "or" joins two thoughts')
    assert by_id["R003"] in (None, "")


def test_dirty_criterion_row_carries_a_finding(workbook):
    ws = workbook["Acceptance Criteria"]
    col = header_row(ws).index("Writing check") + 1
    by_id = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
    # AC003's "then" names two outcomes joined by "and"; AC007 is a single outcome.
    assert 'in the "then" part' in by_id["AC003"]
    assert by_id["AC007"] in (None, "")


def test_summary_sheet_has_writing_check_line(workbook):
    ws = workbook["Summary"]
    items = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert items["Writing check"] == "1 requirement and 1 criterion have warnings; 6 notes"


def test_markdown_puts_findings_in_italics(analysis, recording, tmp_path):
    text = export_markdown(analysis, recording, tmp_path).read_text()
    assert '- *Writing check: warn: "or" joins two thoughts' in text
    assert '  - *Writing check: info: "and" in the "then" part' in text
    assert "- Writing check: 1 requirement and 1 criterion have warnings; 6 notes" in text
