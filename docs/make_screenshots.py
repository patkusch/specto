"""Make the README pictures from one specto run.

Usage: .venv/bin/python docs/make_screenshots.py OUT_DIR

OUT_DIR is the folder a `specto run` wrote (it must hold report.html,
analysis.xlsx and frames/). The pictures land next to this script:

  report-requirements.png   first requirement card with its criteria
  report-flow.png           the screen-flow diagram
  report-questions.png      the SME questions table
  workbook-requirements.png first rows of the Requirements sheet
  workbook-questions.png    first rows of the SME Questions sheet
  frame-example.png         one frame from the recording, 800 wide

Every PNG is kept under 400 KB. Needs Playwright with chromium, openpyxl
and Pillow, all in the project venv.
"""

from __future__ import annotations

import html
import io
import sys
from pathlib import Path

import openpyxl
from PIL import Image
from playwright.sync_api import sync_playwright

DOCS = Path(__file__).resolve().parent
LIMIT = 400 * 1024
VIEWPORT = {"width": 1200, "height": 800}
SHEET_ROWS = 6          # header plus five data rows
FRAME = "frames/frame_0001.jpg"   # the Customer Details screen of the example


# ------------------------------------------------------------------ helpers


def shrink(path: Path) -> None:
    """Bring a PNG under LIMIT: first quantise to 256 colours, then scale down."""
    if path.stat().st_size <= LIMIT:
        return
    im = Image.open(path).convert("RGB")
    im.quantize(colors=256, method=Image.Quantize.MEDIANCUT).save(path, optimize=True)
    while path.stat().st_size > LIMIT and im.width > 800:
        im = im.resize((int(im.width * 0.8), int(im.height * 0.8)), Image.LANCZOS)
        im.quantize(colors=256, method=Image.Quantize.MEDIANCUT).save(path, optimize=True)


def clip_of(page, selector: str, pad: int = 0) -> dict:
    """Page-coordinate box of the first element matching selector, padded."""
    page.evaluate("window.scrollTo(0, 0)")
    box = page.locator(selector).first.bounding_box()
    if box is None:
        raise SystemExit(f"nothing on the page matches {selector}")
    return {"x": max(box["x"] - pad, 0), "y": max(box["y"] - pad, 0),
            "width": box["width"] + 2 * pad, "height": box["height"] + 2 * pad}


def union(a: dict, b: dict) -> dict:
    x = min(a["x"], b["x"])
    y = min(a["y"], b["y"])
    right = max(a["x"] + a["width"], b["x"] + b["width"])
    bottom = max(a["y"] + a["height"], b["y"] + b["height"])
    return {"x": x, "y": y, "width": right - x, "height": bottom - y}


def shoot(page, clip: dict, out: Path) -> None:
    page.screenshot(path=str(out), full_page=True, clip=clip)
    shrink(out)


# ------------------------------------------------------------------ report.html


def report_pictures(page, report: Path) -> None:
    page.goto(report.resolve().as_uri())
    page.wait_for_load_state("networkidle")
    # The requirements picture: section heading, filter box and the first card
    # (the card holds its own acceptance-criteria table).
    heading = clip_of(page, "#requirements h2")
    card = clip_of(page, "#requirements article.req")
    box = union(heading, card)
    box["x"] = 0
    box["width"] = VIEWPORT["width"]
    box["height"] += 12
    shoot(page, box, DOCS / "report-requirements.png")
    shoot(page, clip_of(page, "#flow .flow svg", pad=8), DOCS / "report-flow.png")
    shoot(page, clip_of(page, "#questions table.questions", pad=8), DOCS / "report-questions.png")


# ------------------------------------------------------------------ analysis.xlsx


SHEET_CSS = """
body { margin:16px; background:#fff; font:12px/1.35 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; color:#1f2328; }
table { border-collapse:collapse; table-layout:fixed; }
th, td { border:1px solid #d0d7de; padding:4px 5px; text-align:left; vertical-align:top; overflow-wrap:break-word; }
th.h, td.h { background:#DDEBF7; font-weight:700; }
th.corner, td.rownum, th.col { background:#f0f0f0; color:#555; font-weight:400; text-align:center; width:34px; }
.tab { display:inline-block; margin-top:8px; padding:4px 14px; border:1px solid #d0d7de; border-top:0; background:#fff; font-size:12px; color:#1f2328; }
"""


def col_letter(i: int) -> str:
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def sheet_html(ws, rows: int, keep: int | None = None) -> str:
    """A small HTML table that looks like the top of the sheet.

    Column widths follow the sheet's own (Excel character units, about 6.5
    pixels each); `keep` caps the column count for a very wide sheet.
    """
    ncols = ws.max_column if keep is None else min(ws.max_column, keep)
    widths = []
    for c in range(1, ncols + 1):
        dim = ws.column_dimensions.get(col_letter(c))
        widths.append((dim.width if dim and dim.width else 10) * 6.5)
    parts = [f"<style>{SHEET_CSS}</style><div id='sheet' style='width:{int(sum(widths)) + 40}px'>"
             "<table><colgroup><col style='width:34px'>"]
    parts += [f"<col style='width:{int(w)}px'>" for w in widths]
    parts.append("</colgroup><tr><th class='corner'></th>")
    parts += [f"<th class='col'>{col_letter(c)}</th>" for c in range(1, ncols + 1)]
    parts.append("</tr>")
    for r in range(1, rows + 1):
        parts.append(f"<tr><td class='rownum'>{r}</td>")
        for c in range(1, ncols + 1):
            v = ws.cell(r, c).value
            text = html.escape("" if v is None else str(v))
            cls = "h" if r == 1 else ""
            parts.append(f"<td class='{cls}'>{text}</td>")
        parts.append("</tr>")
    parts.append(f"</table><span class='tab'>{html.escape(ws.title)}</span></div>")
    return "".join(parts)


def workbook_pictures(page, workbook: Path) -> None:
    wb = openpyxl.load_workbook(workbook)
    for sheet, name in (("Requirements", "workbook-requirements.png"),
                        ("SME Questions", "workbook-questions.png")):
        page.set_content(sheet_html(wb[sheet], SHEET_ROWS))
        shoot(page, clip_of(page, "#sheet", pad=16), DOCS / name)


# ------------------------------------------------------------------ frame


def frame_picture(frame: Path) -> None:
    im = Image.open(frame).convert("RGB")
    im = im.resize((800, round(im.height * 800 / im.width)), Image.LANCZOS)
    out = DOCS / "frame-example.png"
    im.save(out, optimize=True)
    shrink(out)


# ------------------------------------------------------------------ main


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2
    out_dir = Path(argv[1])
    for needed in ("report.html", "analysis.xlsx", FRAME):
        if not (out_dir / needed).exists():
            print(f"missing {out_dir / needed}; run specto first")
            return 1
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport=VIEWPORT, device_scale_factor=2)
        report_pictures(page, out_dir / "report.html")
        workbook_pictures(page, out_dir / "analysis.xlsx")
        browser.close()
    frame_picture(out_dir / FRAME)
    for png in sorted(DOCS.glob("*.png")):
        im = Image.open(png)
        print(f"{png.name:28} {im.width}x{im.height}  {png.stat().st_size / 1024:6.0f} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
