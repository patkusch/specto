"""Read the analyst's answers back out of the workbook.

After the follow-up call with the expert, the analyst types what was said into
the Answer and Status columns of the SME Questions sheet. This module reads
those cells, matches them to the questions by id, and stores them in
analysis.json so the next stage (resolve) can turn them into requirements.

Columns are found by their header text, not by position, so a workbook where
someone inserted or moved a column still reads correctly.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from specto.model import Analysis

QUESTIONS_SHEET = "SME Questions"
ID_HEADER = "Id"
ANSWER_HEADER = "Answer"
STATUS_HEADER = "Status"

ANSWERED_WORDS = {"answered", "done", "yes", "closed", "resolved"}
NOT_NEEDED_WORDS = {"not needed", "n/a", "na", "drop", "no longer needed"}

Answers = dict[str, tuple[Optional[str], Optional[str]]]


class ImportResult(BaseModel):
    """What happened when the answers were applied."""

    changed: int = Field(default=0, description="Questions whose answer or status changed")
    unknown_ids: list[str] = Field(default_factory=list, description="Ids in the sheet that no question has")


# ---------------------------------------------------------------- normalising


def normalise_status(text: Optional[str]) -> Optional[str]:
    """Turn what the analyst typed into one of the two closed statuses, or None.

    None means "leave the question's status as it is": the cell was empty, said
    "open", or said something we do not recognise.
    """
    if text is None:
        return None
    word = " ".join(str(text).strip().lower().split())
    if word in ANSWERED_WORDS:
        return "answered"
    if word in NOT_NEEDED_WORDS:
        return "not needed"
    return None


def _cell_text(value: object) -> Optional[str]:
    """The cell's content as stripped text, or None when it is empty."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


# ------------------------------------------------------------------- reading


def _find_columns(ws, wanted: list[str]) -> dict[str, int]:
    """Map each wanted header to its 1-based column number, matching by text."""
    found: dict[str, int] = {}
    for cell in ws[1]:
        text = _cell_text(cell.value)
        if text is None:
            continue
        for name in wanted:
            if name not in found and text.lower() == name.lower():
                found[name] = cell.column
    missing = [name for name in wanted if name not in found]
    if missing:
        raise ValueError(f"Sheet '{ws.title}' has no column headed {', '.join(missing)}")
    return found


def read_answers(xlsx_path: Path | str) -> Answers:
    """Read id -> (answer, status) from the SME Questions sheet.

    Only rows with something typed in Answer or Status are returned. The status
    is normalised (see `normalise_status`); a row with an answer and an empty
    status cell counts as "answered". A row with an answer and "open" typed in
    keeps its status as it is, so a partial answer can be recorded without
    closing the question.
    """
    wb = load_workbook(xlsx_path, read_only=False, data_only=True)
    if QUESTIONS_SHEET not in wb.sheetnames:
        raise ValueError(f"{xlsx_path} has no sheet called '{QUESTIONS_SHEET}'")
    ws = wb[QUESTIONS_SHEET]
    columns = _find_columns(ws, [ID_HEADER, ANSWER_HEADER, STATUS_HEADER])

    answers: Answers = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        def at(name: str) -> Optional[str]:
            index = columns[name] - 1
            return _cell_text(row[index]) if index < len(row) else None

        question_id = at(ID_HEADER)
        answer = at(ANSWER_HEADER)
        status_text = at(STATUS_HEADER)
        if question_id is None or (answer is None and status_text is None):
            continue
        status = normalise_status(status_text)
        if answer is not None and status_text is None:
            status = "answered"
        answers[question_id] = (answer, status)
    return answers


# ------------------------------------------------------------------ applying


def apply_answers(analysis: Analysis, answers: Answers) -> tuple[Analysis, ImportResult]:
    """Write the answers onto the matching questions, in place.

    Returns the analysis and a result saying how many questions changed and
    which ids had no question. A question is counted as changed only when its
    answer or status actually differs, so importing the same sheet twice
    reports zero the second time.
    """
    by_id = {q.id: q for q in analysis.questions}
    result = ImportResult()
    for question_id, (answer, status) in answers.items():
        question = by_id.get(question_id)
        if question is None:
            result.unknown_ids.append(question_id)
            continue
        changed = False
        if answer is not None and answer != question.answer:
            question.answer = answer
            changed = True
        if status is not None and status != question.status:
            question.status = status
            changed = True
        if changed:
            result.changed += 1
    return analysis, result


def import_answers(out_dir: Path | str, xlsx_path: Optional[Path | str] = None) -> ImportResult:
    """Read the workbook's answers into out_dir/analysis.json and save it.

    The workbook defaults to out_dir/analysis.xlsx; pass a path when the
    analyst worked on a copy. This does not re-export the workbook.
    """
    out_dir = Path(out_dir)
    analysis_path = out_dir / "analysis.json"
    analysis = Analysis.model_validate_json(analysis_path.read_text())
    answers = read_answers(Path(xlsx_path) if xlsx_path else out_dir / "analysis.xlsx")
    analysis, result = apply_answers(analysis, answers)
    analysis_path.write_text(analysis.model_dump_json(indent=2))
    return result
